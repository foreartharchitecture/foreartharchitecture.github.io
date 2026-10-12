# Generates invented test files for the local mock Drive (no real firm data).
import os, random, datetime
from openpyxl import Workbook
from openpyxl.styles import Font
from docx import Document
from docx.shared import Inches
from PIL import Image
out = os.path.join(os.path.dirname(__file__), 'fixtures'); os.makedirs(out, exist_ok=True)

# ---- Quarterly finance workbook: Income + Expense trackers, data starting in column B, real date cells ----
wb = Workbook(); ws = wb.active; ws.title = 'Q3 Accounts'
def put(r, c, v, fmt=None, bold=False):
    cell = ws.cell(row=r, column=c, value=v)
    if fmt: cell.number_format = fmt
    if bold: cell.font = Font(bold=True)
r = 2
put(r, 2, 'Quarterly Summary', bold=True); r += 1
put(r, 2, 'Total Income'); put(r, 4, 412000, '#,##0'); r += 2
put(r, 2, 'Income Tracker', bold=True); r += 1
for i, h in enumerate(['Date','Quarter','Project Name','Client Name','Income Category','Invoice No','Payment Stage','Amount','Mode','Received By']): put(r, 2+i, h, bold=True)
r += 1
inc = [(datetime.date(2026,7,4),'Hillside Residence','Arjun Varma','Architecture Fees','INV-011','Advance',60000,'Bank Transfer','Nisha'),
       (datetime.date(2026,7,21),'Cliff Cafe','Meera Das','Interior Fees','INV-012','Stage 2',45000,'UPI','Arjun'),
       (datetime.date(2026,8,9),'Hillside Residence','Arjun Varma','Architecture Fees','INV-013','Stage 2',85000,'Bank Transfer','Nisha'),
       (datetime.date(2026,8,30),'Valley Homestay','Joseph K','Architecture Fees','INV-014','Advance',120000,'Bank Transfer','Arjun'),
       (datetime.date(2026,9,15),'Cliff Cafe','Meera Das','Interior Fees','INV-015','Stage 3',52000,'Cash','Arjun'),
       (datetime.date(2026,9,28),'Rahman Residence','Salma Rahman','Architecture Fees','INV-016','Advance',50000,'Bank Transfer','Nisha')]
for d, p, cl, cat, inv, st, amt, mode, by in inc:
    vals = [d, 'Q3', p, cl, cat, inv, st, amt, mode, by]
    for i, v in enumerate(vals): put(r, 2+i, v, 'dd.mm.yyyy' if i == 0 else ('#,##0' if i == 7 else None))
    r += 1
put(r, 9, 412000, '#,##0'); r += 2
put(r, 2, 'Expense Tracker', bold=True); r += 1
for i, h in enumerate(['Date','Quarter','Expense Head','Category','Project','Vendor / Person','Bill No','Payment Mode','Amount','Paid By']): put(r, 2+i, h, bold=True)
r += 1
exp = [(datetime.date(2026,7,5),'Office rent July','Rent','','Landlord','',  'Bank Transfer',18000,'Nisha'),
       (datetime.date(2026,7,31),'Salaries July','Salary','','Team','', 'Bank Transfer',96000,'Arjun'),
       (datetime.date(2026,8,12),'Site travel','Travel','Hillside Residence','Leena Thomas','', 'Cash',3400,'Arjun'),
       (datetime.date(2026,8,31),'Salaries August','Salary','','Team','', 'Bank Transfer',96000,'Arjun'),
       (datetime.date(2026,9,3),'Plotter ink','Office supplies','','Print House','B-221','UPI',7600,'Nisha'),
       (datetime.date(2026,9,30),'Salaries September','Salary','','Team','', 'Bank Transfer',98000,'Arjun')]
for d, head, cat, proj, ven, bill, mode, amt, by in exp:
    vals = [d, 'Q3', head, cat, proj, ven, bill, mode, amt, by]
    for i, v in enumerate(vals): put(r, 2+i, v, 'dd.mm.yyyy' if i == 0 else ('#,##0' if i == 8 else None))
    r += 1
r += 1; put(r, 2, 'Cash & Bank', bold=True); put(r+1, 2, 'Bank'); put(r+1, 3, 186000, '#,##0')
wb.save(os.path.join(out, 'Q3 2026 Finance.xlsx'))

# ---- Invoice .docx with a large embedded image (tests reading only part of the file) ----
random.seed(7)
img = Image.frombytes('RGB', (700, 700), bytes(random.getrandbits(8) for _ in range(700*700*3)))
img.save(os.path.join(out, 'noise.png'))
def invoice(name, project, client, no, date, due, stage, amount):
    d = Document()
    d.add_picture(os.path.join(out, 'noise.png'), width=Inches(1.2))
    d.add_heading('TAX INVOICE', 1)
    for k, v in [('Project', project), ('Client Name', client), ('Invoice Number', no), ('Date', date), ('Due Date', due)]:
        d.add_paragraph(f'{k} : {v}')
    t = d.add_table(rows=3, cols=2)
    for row, (a, b) in zip(t.rows, [('Description', 'Amount'), (stage, f'{amount:,}'), ('Total', f'{amount:,}')]):
        row.cells[0].text = a; row.cells[1].text = b
    d.add_paragraph('Bank: Example Bank · A/c 123456789012 · IFSC EXMP0001234')
    d.save(os.path.join(out, name))
invoice('Hillside Residence Invoice #2.docx', 'Hillside Residence', 'Mr. Arjun Varma', 'FE/INV/2026/013', '09.08.2026', '24.08.2026', 'Stage 2 - Design development', 85000)
invoice('Valley Homestay Invoice #1.docx', 'Valley Homestay', 'Joseph K', 'FE/INV/2026/014', '30.08.2026', '15.09.2026', 'Advance', 120000)
os.remove(os.path.join(out, 'noise.png'))
for f in sorted(os.listdir(out)): print(f, os.path.getsize(os.path.join(out, f)))
